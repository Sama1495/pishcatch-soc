"""SOC incident reports. Output is plain text inside Office files.

Nothing here fetches a URL, renders HTML, or copies a secret into the file.
"""

from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.time import ensure_utc
from app.db.models import Incident
from app.services.actions import describe_action
from app.services.urls import defang

FOOTER = "Quarantine simulated, no mailbox modified. Rules decide the label; AI only explains."

_URL_RE = re.compile(r"(?i)\bhttps?://[^\s<>'\"]+")
_SECRET_RE = re.compile(
    r"(?i)\b(?:sk-[A-Za-z0-9]{8,}|AIza[0-9A-Za-z_\-]{10,}|ghp_[A-Za-z0-9]{10,}|github_pat_[A-Za-z0-9_]{10,})"
)
_TAG_RE = re.compile(r"<[^>]+>")


def case_code(incident_id: str) -> str:
    """Same short id the board shows, so the file matches the card."""
    hex_chars = "".join(ch for ch in incident_id.lower() if ch in "0123456789abcdef")[:6]
    number = int(hex_chars or "0", 16) % 10000
    return f"PC-{number:04d}"


def is_decided(incident: Incident) -> bool:
    if incident.action_responses or incident.verdict_feedback or incident.decisions:
        return True
    return incident.case_status in {"closed", "quarantine_simulated"}


def safe_text(value: object, limit: int = 4000) -> str:
    """Defang links and strip secrets and tags before anything is written."""
    if value is None:
        return ""
    text = str(value)
    text = _SECRET_RE.sub("[redacted-secret]", text)
    text = _URL_RE.sub(lambda match: defang(match.group(0)), text)
    text = _TAG_RE.sub(" ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:limit]


def _stamp(value: datetime | None) -> str:
    if value is None:
        return ""
    return ensure_utc(value).strftime("%Y-%m-%d %H:%M:%S UTC")


def _rows(incident: Incident) -> dict[str, list[list[str]]]:
    headers = incident.headers_json or {}
    policy = incident.policy_json or {}
    actions = list(incident.action_responses or [])
    verdicts = list(incident.verdict_feedback or [])
    decisions = list(incident.decisions or [])
    latest_action = actions[-1] if actions else None
    if latest_action is not None:
        _simulated, action_message = describe_action(latest_action.decision, latest_action.offered_action)
        human = (
            f"{latest_action.decision} by {latest_action.analyst_id} at {_stamp(latest_action.created_at)}. "
            f"{action_message}"
        )
    elif verdicts:
        latest = verdicts[-1]
        human = f"Agent was {latest.verdict}, recorded by {latest.analyst_id} at {_stamp(latest.created_at)}."
    elif decisions:
        latest_decision = decisions[-1]
        human = f"{latest_decision.action} by {latest_decision.analyst_id} at {_stamp(latest_decision.created_at)}."
    else:
        human = incident.case_status.replace("_", " ")

    auth_bits = [
        f"trust {headers.get('auth_trust') or 'absent'}",
        f"SPF {incident.spf_result or 'not scored'}",
        f"DKIM {incident.dkim_result or 'not scored'}",
        f"DMARC {incident.dmarc_result or 'not scored'}",
    ]
    auth_results = headers.get("authentication_results") or []
    if auth_results:
        auth_bits.append("Authentication-Results: " + " | ".join(str(item) for item in auth_results))
    if incident.caveat:
        auth_bits.append(incident.caveat)

    summary = [
        ["Case", case_code(incident.id)],
        ["Incident id", incident.id],
        ["Verdict", incident.risk_label.replace("_", " ")],
        ["Risk score", str(incident.risk_score)],
        ["Confidence", f"{incident.confidence_score or 0}% {incident.confidence_band or 'low'}"],
        ["Human decision", human],
        ["Sender", incident.sender or ""],
        ["Sender address", incident.sender_address or ""],
        ["Reply-To", incident.reply_to or ""],
        ["Subject", incident.subject or ""],
        ["Received", _stamp(incident.created_at)],
        ["Authentication", " ".join(auth_bits)],
        ["Offered action", str(policy.get("offered_action_label") or policy.get("offered_action") or "")],
        ["Policy summary", str(policy.get("summary") or "")],
        ["Plain-text excerpt", safe_text(incident.body_text, 2000)],
    ]
    evidence = [
        [row.rule_id, row.reason, str(row.weight)]
        for row in incident.evidence
    ]
    indicators = [
        [row.indicator_type, row.display_value]
        for row in incident.indicators
        if row.indicator_type in {"url", "domain", "ip", "attachment"}
    ]
    timeline = [
        [row.title, row.status, _stamp(row.started_at), str(row.duration_ms)]
        for row in incident.operation_steps
    ]
    audit: list[list[str]] = []
    for row in actions:
        audit.append([_stamp(row.created_at), row.analyst_id, f"action {row.decision}", row.reason])
    for row in verdicts:
        audit.append([_stamp(row.created_at), row.analyst_id, f"agent {row.verdict}", row.reason])
    for row in decisions:
        audit.append([_stamp(row.created_at), row.analyst_id, row.action, row.reason])
    for row in incident.audit_events:
        audit.append([_stamp(row.created_at), row.actor, row.event_type, ""])
    steps = list((incident.brief_json or {}).get("recommended_actions") or [])
    return {
        "summary": [[safe_text(left, 80), safe_text(right)] for left, right in summary],
        "evidence": [[safe_text(rule, 80), safe_text(reason), safe_text(weight, 16)] for rule, reason, weight in evidence],
        "indicators": [[safe_text(kind, 32), safe_text(value)] for kind, value in indicators],
        "timeline": [[safe_text(a, 120), safe_text(b, 32), safe_text(c, 40), safe_text(d, 16)] for a, b, c, d in timeline],
        "audit": [[safe_text(a, 40), safe_text(b, 80), safe_text(c, 80), safe_text(d)] for a, b, c, d in audit],
        "steps": [safe_text(item) for item in steps],
        "human": safe_text(human),
    }


def build_docx(incident: Incident) -> bytes:
    rows = _rows(incident)
    code = case_code(incident.id)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    title = document.add_heading("PhishCatch SOC incident report", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x0B, 0x3A, 0x4A)
    document.add_paragraph(f"{code}  ·  {incident.id}")
    document.add_heading("Executive summary", level=1)
    summary_map = dict(rows["summary"])
    for label in ("Verdict", "Risk score", "Confidence", "Human decision"):
        document.add_paragraph(f"{label}: {summary_map.get(label, '')}")
    document.add_heading("What happened", level=1)
    for label in ("Sender", "Reply-To", "Subject", "Received", "Authentication"):
        document.add_paragraph(f"{label}: {summary_map.get(label, '')}")
    document.add_heading("Evidence", level=1)
    _docx_table(document, ["Rule", "Reason", "Weight"], rows["evidence"])
    document.add_heading("Defanged indicators", level=1)
    _docx_table(document, ["Type", "Value"], rows["indicators"] or [["", "No indicators were stored."]])
    document.add_heading("Agent timeline", level=1)
    _docx_table(document, ["Step", "Status", "Started", "Duration ms"], rows["timeline"])
    document.add_heading("Policy and the human decision", level=1)
    document.add_paragraph(f"Offered action: {summary_map.get('Offered action', '')}")
    document.add_paragraph(summary_map.get("Policy summary", ""))
    document.add_paragraph(rows["human"])
    document.add_heading("Analyst feedback", level=1)
    feedback = [item for item in rows["audit"] if item[2].startswith("agent ") or item[2] in {"confirm", "correct"}]
    if feedback:
        _docx_table(document, ["When", "Analyst", "Record", "Reason"], feedback)
    else:
        document.add_paragraph("No correct/wrong verdict was recorded.")
    document.add_heading("Recommended next steps", level=1)
    for step in rows["steps"] or ["Review the case with the stored evidence."]:
        document.add_paragraph(step, style="List Bullet")
    document.add_heading("Plain-text excerpt", level=1)
    document.add_paragraph(summary_map.get("Plain-text excerpt", "") or "No text body was stored.")
    footer = document.add_paragraph(FOOTER)
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx(incident: Incident) -> bytes:
    rows = _rows(incident)
    book = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0B3A4A")
    _sheet(book.active, "Summary", ["Field", "Value"], rows["summary"], header_font, header_fill)
    _sheet(book.create_sheet("Evidence"), "Evidence", ["Rule", "Reason", "Weight"], rows["evidence"], header_font, header_fill)
    _sheet(book.create_sheet("Indicators"), "Indicators", ["Type", "Defanged value"], rows["indicators"], header_font, header_fill)
    _sheet(book.create_sheet("Timeline"), "Timeline", ["Step", "Status", "Started", "Duration ms"], rows["timeline"], header_font, header_fill)
    audit_rows = rows["audit"] + [["", "", "footer", FOOTER]]
    _sheet(book.create_sheet("Audit"), "Audit", ["When", "Who", "Event", "Detail"], audit_rows, header_font, header_fill)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _docx_table(document: Document, headers: list[str], data: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for record in data:
        cells = table.add_row().cells
        for index, value in enumerate(record):
            cells[index].text = value
    document.add_paragraph("")


def _sheet(sheet, title: str, headers: list[str], data: list[list[str]], header_font: Font, header_fill: PatternFill) -> None:
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    for record in data:
        sheet.append(record)
    for index, _header in enumerate(headers, start=1):
        width = 18 if index == 1 else 42
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.left.text = "PhishCatch SOC"
    sheet.oddFooter.left.text = FOOTER
    sheet.page_margins.left = 0.5
    sheet.page_margins.right = 0.5
