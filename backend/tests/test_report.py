"""Incident reports are Office zip files and never carry live links or secrets."""

import io

from docx import Document
from openpyxl import load_workbook

from app.services.report import FOOTER

RAW = b"""From: Alert <alerts@phish.example>
Reply-To: reply@phish.example
To: analyst@contoso.example
Subject: Report fixture
Date: Tue, 1 Oct 2024 12:00:00 +0000
MIME-Version: 1.0
Content-Type: text/plain; charset=utf-8

Please open https://phish.example/login and use sk-FAKEKEY1234567890
"""


def _upload(client):
    return client.post("/api/v1/incidents", files={"file": ("report-fixture.eml", RAW, "message/rfc822")})


def _decide(client, incident_id: str) -> None:
    feedback = client.post(
        f"/api/v1/incidents/{incident_id}/feedback",
        json={"verdict": "correct", "reason": "Matches the evidence.", "analyst_id": "analyst.1"},
    )
    assert feedback.status_code == 200, feedback.text
    action = client.post(
        f"/api/v1/incidents/{incident_id}/actions/response",
        json={"decision": "approve", "reason": "Approved the offered action.", "analyst_id": "analyst.1"},
    )
    assert action.status_code == 200, action.text


def _docx_text(payload: bytes) -> str:
    document = Document(io.BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _xlsx_text(payload: bytes) -> str:
    book = load_workbook(io.BytesIO(payload))
    parts = [" ".join(book.sheetnames)]
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            parts.append(" ".join("" if value is None else str(value) for value in row))
    return "\n".join(parts)


def test_reports_are_zip_files_and_defang_body_secrets(client):
    created = _upload(client)
    assert created.status_code == 201, created.text
    incident_id = created.json()["incident_id"]
    blocked = client.get(f"/api/v1/incidents/{incident_id}/report.docx")
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "report_not_ready"

    _decide(client, incident_id)
    docx = client.get(f"/api/v1/incidents/{incident_id}/report.docx")
    xlsx = client.get(f"/api/v1/incidents/{incident_id}/report.xlsx")
    assert docx.status_code == 200, docx.text
    assert xlsx.status_code == 200, xlsx.text
    assert docx.content[:2] == b"PK"
    assert xlsx.content[:2] == b"PK"

    doc_text = _docx_text(docx.content)
    sheet_text = _xlsx_text(xlsx.content)
    for text in (doc_text, sheet_text):
        assert "https://phish.example/login" not in text
        assert "sk-FAKEKEY1234567890" not in text
        assert "phish[.]example" in text
        assert "[redacted-secret]" in text
        assert FOOTER in text
        assert "<a " not in text.lower()

    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Summary", "Evidence", "Indicators", "Timeline", "Audit"]
    assert "PC-" in doc_text
